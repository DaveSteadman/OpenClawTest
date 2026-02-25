#!/usr/bin/env python3

import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path

SKILLS_ROOT = Path(__file__).resolve().parent.parent
if str(SKILLS_ROOT) not in sys.path:
    sys.path.insert(0, str(SKILLS_ROOT))

from CommonCode.FolderNavigator import FolderNavigator

try:
    from docx import Document
except ImportError:
    print(json.dumps({"status": "error", "error": "Missing dependency python-docx. Install with: pip install -r requirements.txt"}, ensure_ascii=False))
    raise SystemExit(1)


def validate_domain(domain: str) -> str:
    value = (domain or "").strip()
    if not value:
        raise ValueError("Domain must be a non-empty string")
    if not re.fullmatch(r"[A-Za-z]+", value):
        raise ValueError("Domain must be alphabetic only (A-Z, a-z)")
    return value


def parse_timeframe(value: str) -> dt.date:
    cleaned = (value or "").strip().replace("-", "/")
    parts = cleaned.split("/")
    if len(parts) != 3:
        raise ValueError("--timeframe must be YYYY/MM/DD or YYYY-MM-DD")
    year, month, day = (int(part) for part in parts)
    return dt.date(year, month, day)


def safe_name(text: str, fallback: str = "Presentation") -> str:
    tokens = re.findall(r"[A-Za-z0-9]+", text or "")
    if not tokens:
        return fallback
    output = []
    for token in tokens:
        if token.isdigit():
            output.append(token)
        else:
            output.append(token[0].upper() + token[1:].lower())
    return "".join(output)[:80]


def resolve_analysis_day_dir(navigator: FolderNavigator, source_domain: str, as_of_date: dt.date) -> Path:
    root = navigator.get_domain_root("analyse", source_domain)
    day_dir = root / f"{as_of_date.year:04d}" / f"{as_of_date.month:02d}" / f"{as_of_date.day:02d}"
    if not day_dir.exists():
        raise FileNotFoundError(f"No analysis day folder found: {day_dir}")
    return day_dir


def parse_markdown_sections(markdown_text: str) -> list[tuple[str, list[str]]]:
    sections: list[tuple[str, list[str]]] = []
    current_title = "Overview"
    current_lines: list[str] = []

    for raw_line in markdown_text.splitlines():
        line = raw_line.strip()
        if line.startswith("## "):
            if current_lines:
                sections.append((current_title, current_lines))
            current_title = line[3:].strip() or "Section"
            current_lines = []
            continue
        if line.startswith("# "):
            continue
        if line:
            current_lines.append(line)

    if current_lines:
        sections.append((current_title, current_lines))

    if not sections:
        sections.append(("Overview", ["No sectioned content available."]))

    return sections


def build_docx(title: str, source_domain: str, date_value: dt.date, source_files: list[Path], sections: list[tuple[str, list[str]]], out_path: Path) -> None:
    document = Document()
    document.add_heading(title, level=0)
    document.add_paragraph(f"Domain: {source_domain}")
    document.add_paragraph(f"Date: {date_value.isoformat()}")

    for section_title, lines in sections:
        document.add_heading(section_title, level=1)
        for line in lines:
            if line.startswith("-"):
                document.add_paragraph(line.lstrip("-").strip(), style="List Bullet")
            else:
                document.add_paragraph(line)

    document.add_heading("Source Files", level=1)
    for source in source_files:
        document.add_paragraph(source.name, style="List Bullet")

    document.save(str(out_path))


def main() -> int:
    parser = argparse.ArgumentParser(description="Build a manager-ready DOCX presentation from analysis markdown outputs")
    parser.add_argument("domain", help="Output domain (alphabetic only)")
    parser.add_argument("--timeframe", required=True, help="YYYY/MM/DD or YYYY-MM-DD")
    parser.add_argument("--source-domain", default="", help="Optional source analysis domain, defaults to domain")
    parser.add_argument("--title", default="", help="Optional presentation title")
    parser.add_argument("--max-files", type=int, default=4, help="Max analysis markdown files to include")

    try:
        args = parser.parse_args()
        output_domain = validate_domain(args.domain)
        source_domain = validate_domain((args.source_domain or "").strip() or output_domain)
        date_value = parse_timeframe(args.timeframe)
    except Exception as exc:
        print(json.dumps({"status": "error", "error": str(exc)}, ensure_ascii=False))
        return 1

    navigator = FolderNavigator.from_fixed_point()

    try:
        day_dir = resolve_analysis_day_dir(navigator, source_domain, date_value)
        analysis_files = sorted(day_dir.glob("*.md"))[: max(1, args.max_files)]
        if not analysis_files:
            raise FileNotFoundError(f"No .md analysis files in {day_dir}")

        merged = "\n\n".join(path.read_text(encoding="utf-8", errors="replace") for path in analysis_files)
        sections = parse_markdown_sections(merged)

        present_root = navigator.get_domain_root("present", output_domain, create=True)
        out_dir = present_root / f"{date_value.year:04d}" / f"{date_value.month:02d}" / f"{date_value.day:02d}"
        out_dir.mkdir(parents=True, exist_ok=True)

        title = (args.title or "").strip() or f"{source_domain} Executive Brief"
        out_path = out_dir / f"{safe_name(title, fallback='ExecutiveBrief')}_{date_value.isoformat()}.docx"

        build_docx(title, source_domain, date_value, analysis_files, sections, out_path)

        print(json.dumps({
            "status": "ok",
            "domain": output_domain,
            "source_domain": source_domain,
            "timeframe": date_value.strftime("%Y/%m/%d"),
            "analysis_files": [str(path) for path in analysis_files],
            "output_path": str(out_path),
        }, ensure_ascii=False))
        return 0
    except Exception as exc:
        print(json.dumps({
            "status": "error",
            "domain": output_domain,
            "source_domain": source_domain,
            "timeframe": date_value.strftime("%Y/%m/%d"),
            "error": str(exc),
        }, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
